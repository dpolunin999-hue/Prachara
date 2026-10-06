from __future__ import annotations

import argparse
import re
import sys
from datetime import date
from pathlib import Path
from urllib.parse import quote, urljoin, urlsplit, urlunsplit

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT = Path(__file__).resolve().parent.parent
DOCS = PROJECT / "docs"
CONFIG = PROJECT / "mkdocs.yml"
OUTPUT_DIR = PROJECT / "Документы_системы"
OUTPUT = OUTPUT_DIR / "Система_Прачары.docx"
SITE_ROOT = "https://dpolunin999-hue.github.io/Prachara/"
TODAY = date.today().strftime("%d.%m.%Y")
FONT = "Arial"
DARK_GREEN = "2E5D50"
PALE_GRAY = "F5F5F3"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, name=FONT, size=None, bold=None, italic=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr()
    fonts = run._element.rPr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run._element.rPr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def keep_table_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def add_hyperlink(paragraph, text: str, url: str, bold=False, italic=False):
    relation_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5E4A")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    if italic:
        r_pr.append(OxmlElement("w:i"))
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)
    r_pr.append(fonts)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def clean_text(text: str) -> str:
    text = re.sub(r"\{[^{}]*\}\s*$", "", text)
    text = re.sub(r"<[^>]+>", "", text)
    return text.replace("&nbsp;", " ").strip()


INLINE = re.compile(
    r"(\*\*.+?\*\*|\*[^*]+\*|\x60[^\x60]+\x60|!\[[^\]]*\]\([^)]+\)|\[[^\]]+\]\([^)]+\))"
)


def add_inline(paragraph, text: str, base_bold=False, base_italic=False) -> None:
    text = clean_text(text)
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_font(run, bold=base_bold, italic=base_italic)
        token = match.group(0)
        if token.startswith("**"):
            inner = token[2:-2]
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", inner)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2), bold=True, italic=base_italic)
            else:
                run = paragraph.add_run(inner)
                set_font(run, bold=True, italic=base_italic)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, bold=base_bold, italic=True)
        elif token.startswith(chr(96)):
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Consolas", size=9, color="333333")
        elif token.startswith("!["):
            image = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", token)
            if image:
                run = paragraph.add_run(f"[Изображение: {image.group(1) or 'без подписи'}]")
                set_font(run, italic=True, color="666666")
        else:
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1).replace("**", ""), link.group(2))
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_font(run, bold=base_bold, italic=base_italic)


def parse_nav(config_text: str):
    items = []
    current = None
    in_nav = False
    pattern = re.compile(r"^(\s*)-\s+([^:]+):\s*(.*)$")
    for raw in config_text.splitlines():
        if raw.strip() == "nav:":
            in_nav = True
            continue
        if not in_nav:
            continue
        if raw and not raw.startswith(" "):
            break
        match = pattern.match(raw)
        if not match:
            continue
        indent = len(match.group(1))
        label = match.group(2).strip()
        value = match.group(3).strip()
        if indent == 2:
            current = {"label": label, "path": value if value.endswith(".md") else None, "children": []}
            items.append(current)
        elif indent >= 6 and current is not None and value.endswith(".md"):
            current["children"].append({"label": label, "path": value})
    return items


def site_url_for(path: str) -> str:
    source = Path(path)
    parts = [quote(part) for part in source.with_suffix("").parts]
    if parts and parts[-1].lower() == "index":
        parts = parts[:-1]
    return SITE_ROOT + "/".join(parts) + ("/" if parts else "")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    title = doc.styles["Title"]
    title.font.name = FONT
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for level, size in {1: 21, 2: 15.5, 3: 12.5, 4: 11}.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(14 if level == 1 else 9)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    for name, left, font_name, size in (
        ("Bullet Text", 0.55, FONT, 10.5),
        ("Checklist Text", 0.55, FONT, 10.5),
        ("Code Text", 0.55, "Consolas", 8.5),
    ):
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = font_name
        style.font.size = Pt(size)
        style.paragraph_format.left_indent = Cm(left)
        style.paragraph_format.first_line_indent = Cm(-0.35 if name != "Code Text" else 0)
        style.paragraph_format.space_after = Pt(3)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Система Прачары   •   ")
    set_font(run, size=8.5, color="666666")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = paragraph.add_run()
    field_run._r.append(begin)
    field_run._r.append(instruction)
    field_run._r.append(separate)
    field_run.add_text("1")
    end_run = paragraph.add_run()
    end_run._r.append(end)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph("")
    paragraph = doc.add_paragraph(style="Title")
    paragraph.add_run("Система Прачары")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Практическое руководство для Центра, секретарей и команд юнитов")
    set_font(run, size=14, color="333333")
    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(20)
    intro.paragraph_format.space_after = Pt(8)
    run = intro.add_run(
        "Документ объединяет актуальные материалы основного сайта в одну последовательную систему. "
        "Он помогает увидеть архитектуру Прачары целиком и быстро перейти от цели к роли, действию, "
        "шаблону и ожидаемому результату."
    )
    set_font(run, size=11.5)
    paragraph = doc.add_paragraph()
    add_hyperlink(paragraph, "Открыть сайт Система Прачары", SITE_ROOT)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(28)
    run = meta.add_run(f"Дата сборки: {TODAY}")
    set_font(run, size=10, color="666666")
    doc.add_page_break()


def add_contents(doc: Document, nav) -> None:
    doc.add_paragraph("Содержание", style="Heading 1")
    for item in nav:
        if item["label"] == "Главная":
            continue
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item["label"])
        set_font(run, size=11, bold=True)
        for child in item["children"]:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.65)
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(child["label"])
            set_font(run, size=10)


def add_table(doc: Document, rows):
    if not rows:
        return
    headers, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    keep_table_row_together(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, DARK_GREEN)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(headers) > 2 and index == 0 else WD_ALIGN_PARAGRAPH.LEFT
        add_inline(paragraph, text, base_bold=True)
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
    for row_index, row in enumerate(body):
        table_row_obj = table.add_row()
        keep_table_row_together(table_row_obj)
        cells = table_row_obj.cells
        for index, text in enumerate(row):
            cell = cells[index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, PALE_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 and len(text) < 14 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(9.2)
    doc.add_paragraph()


def table_row(line: str):
    return [clean_text(part) for part in line.strip().strip("|").split("|")]


def is_separator(row):
    return all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in row)


def add_markdown(doc: Document, path: Path, heading_base: int = 1) -> None:
    first_paragraph = len(doc.paragraphs)
    content = path.read_text(encoding="utf-8-sig")
    content = re.sub(r"\A---\s*\n.*?\n---\s*\n", "", content, count=1, flags=re.S)
    # Body links in Word open the matching public site page.
    source_url = SITE_ROOT + path.relative_to(DOCS).as_posix()
    def public_link(match):
        target = urljoin(source_url, match.group(2))
        parts = urlsplit(target)
        target_path = parts.path[:-3] + "/" if parts.path.endswith(".md") else parts.path
        return match.group(1) + urlunsplit((parts.scheme, parts.netloc, target_path, parts.query, parts.fragment)) + match.group(3)
    content = re.sub(r"(\]\()([^\s)]+)(\))", public_link, content)
    lines = content.splitlines()
    index = 0
    skipped_h1 = False
    in_code = False
    code_kind = ""
    code_lines = []
    callout_indent = False

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if stripped.startswith(chr(96) * 3):
            if not in_code:
                in_code = True
                code_kind = stripped.lstrip(chr(96)).strip()
                code_lines = []
            else:
                if code_kind == "mermaid":
                    paragraph = doc.add_paragraph()
                    add_inline(paragraph, "Схема доступна в интерактивной версии сайта.", base_italic=True)
                else:
                    for code_line in code_lines:
                        paragraph = doc.add_paragraph(style="Code Text")
                        for part in re.split(r"([\U0001F300-\U0001FAFF\u2600-\u27BF\uFE0F]+)", code_line):
                            if not part:
                                continue
                            run = paragraph.add_run(part)
                            emoji = bool(re.fullmatch(r"[\U0001F300-\U0001FAFF\u2600-\u27BF\uFE0F]+", part))
                            set_font(run, name="Segoe UI Emoji" if emoji else "Consolas", size=8.5)
                in_code = False
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|") and lines[index].strip().endswith("|"):
                table_lines.append(table_row(lines[index]))
                index += 1
            add_table(doc, [row for row in table_lines if not is_separator(row)])
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            source_level = len(heading.group(1))
            title = clean_text(heading.group(2))
            if source_level == 1 and not skipped_h1:
                skipped_h1 = True
                index += 1
                continue
            level = min(4, heading_base + max(1, source_level - 1))
            doc.add_paragraph(title, style=f"Heading {level}")
            index += 1
            continue

        callout = re.match(r"^(?:!!!|\?\?\?)\s+\w+\s+[\"'](.+)[\"']$", stripped)
        if callout:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(callout.group(1))
            set_font(run, bold=True)
            callout_indent = True
            index += 1
            continue

        if not stripped or stripped == "---":
            callout_indent = False
            index += 1
            continue

        if raw.startswith("    ") and callout_indent:
            raw = raw[4:]
            stripped = raw.strip()

        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.65)
            add_inline(paragraph, stripped.lstrip("> ").strip(), base_italic=True)
            index += 1
            continue

        check = re.match(r"^-\s+\[([ xX])\]\s+(.+)$", stripped)
        if check:
            paragraph = doc.add_paragraph(style="Checklist Text")
            mark = "☒" if check.group(1).lower() == "x" else "☐"
            add_inline(paragraph, f"{mark} {check.group(2)}")
            index += 1
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet:
            paragraph = doc.add_paragraph(style="Bullet Text")
            add_inline(paragraph, f"• {bullet.group(1)}")
            index += 1
            continue

        ordered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if ordered:
            paragraph = doc.add_paragraph(style="Bullet Text")
            add_inline(paragraph, f"{ordered.group(1)}. {ordered.group(2)}")
            index += 1
            continue

        # Native site disclosures become readable headings in the handbook.
        summary = re.match(r"<summary>(.*?)</summary>", stripped)
        if summary:
            strong = re.search(r"<strong>(.*?)</strong>", summary.group(1))
            title = clean_text(strong.group(1) if strong else summary.group(1))
            if title:
                doc.add_paragraph(title, style=f"Heading {min(4, heading_base + 2)}")
            index += 1
            continue

        if stripped.startswith("<") and stripped.endswith(">"):
            index += 1
            continue

        paragraph = doc.add_paragraph()
        add_inline(paragraph, stripped)
        index += 1

    # Keep the closing status block compact enough that its source link does not
    # become an orphan on a separate page in the generated human edition.
    if path.name == "Правила_общения_с_аудиторией.md":
        for closing_paragraph in doc.paragraphs[-8:]:
            closing_paragraph.paragraph_format.space_after = Pt(0)
            closing_paragraph.paragraph_format.line_spacing = 1.0

    if path.name == "Контент_мейкер.md":
        for role_paragraph in doc.paragraphs[first_paragraph:]:
            role_paragraph.paragraph_format.space_after = Pt(2)
            role_paragraph.paragraph_format.line_spacing = 1.05
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    relative = str(path.relative_to(DOCS)).replace("\\", "/")
    add_hyperlink(paragraph, "Версия страницы на сайте", site_url_for(relative))


def build() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    nav = parse_nav(CONFIG.read_text(encoding="utf-8-sig"))
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_contents(doc, nav)

    first_section = True
    for item in nav:
        if item["label"] == "Главная":
            continue
        heading = doc.add_paragraph(item["label"], style="Heading 1")
        heading.paragraph_format.page_break_before = True
        first_section = False
        if item["path"]:
            add_markdown(doc, DOCS / item["path"], heading_base=1)
        else:
            for child in item["children"]:
                doc.add_paragraph(child["label"], style="Heading 2")
                add_markdown(doc, DOCS / child["path"], heading_base=2)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    props = doc.core_properties
    props.title = "Система Прачары"
    props.subject = "Практическое руководство по системе Прачары"
    props.author = "Департамент Прачары"
    props.keywords = "Прачара, юнит, забота, привлечение, программы, обучение"

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Собрать Word-издание системы Прачары")
    parser.add_argument("--build-date", type=date.fromisoformat, help="Дата сборки YYYY-MM-DD; по умолчанию дата компьютера")
    args = parser.parse_args()
    if args.build_date:
        TODAY = args.build_date.strftime("%d.%m.%Y")
    try:
        print(build())
    except Exception as exc:
        print(f"Ошибка сборки: {exc}", file=sys.stderr)
        raise
